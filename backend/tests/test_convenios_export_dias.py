"""Exportação de Convênios Estaduais: coluna "Dias p/ fim da vigência".

Pedido do dono (10/09/2026): na exportação (PDF/Excel/Word) da tela dos
ESTADUAIS, uma coluna com os dias que faltam para o fim da vigência.

A coluna entra nos TRÊS formatos (regra do módulo: "uma coluna nova entra nos
três ou em nenhum") e usa a MESMA data de fim de vigência que a coluna "Fim da
vigência" já usa (`dt_vigencia_atual or dt_vigencia_final`) — negativo = já
venceu, vazio = sem data. Número no Excel (ordenável), texto no PDF/Word.
"""
import io
from datetime import date, timedelta

from services import convenios_export as X


class _Conv:
    """ConvenioEstadual mínimo para `linha_de` — só os atributos que ela lê."""
    def __init__(self, venc_atual=None, venc_final=None):
        self.raw_data = {}
        self.nr_plano_trabalho = None
        self.nr_sigcon = "123/2026"
        self.fonte = "SIGCON-MG"
        self.orgao_concedente = "SEGOV"
        self.objetivo = None
        self.objeto = "Objeto teste"
        self.situacao = "Em vigor"
        self.valor_concedente = 100.0
        self.valor_total = 100.0
        self.dt_vigencia_inicial = date(2026, 1, 1)
        self.dt_vigencia_atual = venc_atual
        self.dt_vigencia_final = venc_final


def test_coluna_existe_nos_tres_formatos_e_a_linha_bate():
    chaves = [c[0] for c in X.COLUNAS]
    assert "dias_vigencia" in chaves
    # invariante do módulo: a linha tem exatamente as chaves das colunas
    assert set(X.linha_de(_Conv(venc_atual=date(2026, 12, 31)))) == set(chaves)


def test_dias_positivo_negativo_e_ausente():
    hoje = date.today()
    # faltam 10 dias
    assert X.linha_de(_Conv(venc_atual=hoje + timedelta(days=10)))["dias_vigencia"] == 10
    # venceu há 5 dias -> negativo
    assert X.linha_de(_Conv(venc_atual=hoje - timedelta(days=5)))["dias_vigencia"] == -5
    # sem data de vigência -> None (não inventa prazo)
    assert X.linha_de(_Conv())["dias_vigencia"] is None


def test_prioridade_vigencia_atual_sobre_final():
    """A vigência ATUAL (com aditivos) manda; a final é fallback — igual à tela e
    à coluna 'Fim da vigência', para as duas colunas nunca divergirem."""
    hoje = date.today()
    l = X.linha_de(_Conv(venc_atual=hoje + timedelta(days=3),
                         venc_final=hoje + timedelta(days=999)))
    assert l["dias_vigencia"] == 3
    assert l["vigencia"] == hoje + timedelta(days=3)


def test_texto_pdf_word():
    assert X.dias_vigencia_txt(None) == "-"
    assert X.dias_vigencia_txt(0) == "0"      # vence hoje: não some
    assert X.dias_vigencia_txt(12) == "12"
    assert X.dias_vigencia_txt(-7) == "-7"


def test_xlsx_tem_cabecalho_e_valor_numerico():
    from openpyxl import load_workbook
    hoje = date.today()
    linhas = [X.linha_de(_Conv(venc_atual=hoje + timedelta(days=20)))]
    buf = io.BytesIO(X.gerar_xlsx(linhas, titulo="T", recorte=[],
                                  emitido_em=__import__("datetime").datetime.now()))
    ws = load_workbook(buf)["Convênios"]
    cabecalhos = [c.value for c in ws[1]]
    assert "Dias p/ fim da vigência" in cabecalhos
    col = cabecalhos.index("Dias p/ fim da vigência") + 1
    cel = ws.cell(2, col)
    assert cel.value == 20            # número cru (ordenável), não texto
    assert cel.number_format == "0"


def test_docx_mostra_a_coluna():
    from docx import Document
    hoje = date.today()
    linhas = [X.linha_de(_Conv(venc_atual=hoje + timedelta(days=8)))]
    buf = io.BytesIO(X.gerar_docx(linhas, titulo="T", subtitulo="s", recorte=[],
                                  emitido_em=__import__("datetime").datetime.now()))
    doc = Document(buf)
    tab = doc.tables[0]
    cabecalhos = [c.text for c in tab.rows[0].cells]
    assert "Dias p/ fim da vigência" in cabecalhos
    col = cabecalhos.index("Dias p/ fim da vigência")
    assert tab.rows[1].cells[col].text == "8"


# --- data de pagamento + data de empenho (fonte: Transparência-MG, como o RM) ---

def test_pagamento_empenho_entram_nas_colunas_e_na_linha():
    chaves = [c[0] for c in X.COLUNAS]
    assert "dt_pagamento" in chaves and "dt_empenho" in chaves
    mg = {"dt_pagamento": date(2026, 4, 8), "dt_empenho": date(2026, 2, 1)}
    linha = X.linha_de(_Conv(venc_atual=date(2026, 12, 31)), mg=mg)
    # a invariante do módulo continua valendo COM o mg
    assert set(linha) == set(chaves)
    assert linha["dt_pagamento"] == date(2026, 4, 8)
    assert linha["dt_empenho"] == date(2026, 2, 1)


def test_sem_mg_as_colunas_ficam_vazias_nao_inventadas():
    """Convênio sem dado de MG (ou tenant não-MG): as duas colunas são None e
    saem '-'/em branco — nunca uma data inventada."""
    from services.rm_pdf import _fmt_dt
    linha = X.linha_de(_Conv(venc_atual=date(2026, 12, 31)))  # sem mg
    assert linha["dt_pagamento"] is None and linha["dt_empenho"] is None
    assert _fmt_dt(linha["dt_pagamento"]) == ""      # render vazio, não "-inventado-"


def test_pagamento_empenho_no_xlsx():
    from openpyxl import load_workbook
    mg = {"dt_pagamento": date(2026, 4, 8), "dt_empenho": date(2026, 2, 1)}
    linhas = [X.linha_de(_Conv(venc_atual=date(2026, 12, 31)), mg=mg)]
    buf = io.BytesIO(X.gerar_xlsx(linhas, titulo="T", recorte=[],
                                  emitido_em=__import__("datetime").datetime.now()))
    ws = load_workbook(buf)["Convênios"]
    cab = [c.value for c in ws[1]]
    assert "Data de pagamento" in cab and "Data de empenho" in cab
    cp = ws.cell(2, cab.index("Data de pagamento") + 1)
    ce = ws.cell(2, cab.index("Data de empenho") + 1)
    # openpyxl lê data de volta como datetime — comparo só a parte de data.
    def _d(v):
        return v.date() if hasattr(v, "date") else v
    assert _d(cp.value) == date(2026, 4, 8) and cp.number_format == "DD/MM/YYYY"
    assert _d(ce.value) == date(2026, 2, 1)


def test_pagamento_empenho_no_docx():
    from docx import Document
    mg = {"dt_pagamento": date(2026, 4, 8), "dt_empenho": date(2026, 2, 1)}
    linhas = [X.linha_de(_Conv(venc_atual=date(2026, 12, 31)), mg=mg)]
    buf = io.BytesIO(X.gerar_docx(linhas, titulo="T", subtitulo="s", recorte=[],
                                  emitido_em=__import__("datetime").datetime.now()))
    tab = Document(buf).tables[0]
    cab = [c.text for c in tab.rows[0].cells]
    assert "Data de pagamento" in cab and "Data de empenho" in cab
    assert tab.rows[1].cells[cab.index("Data de pagamento")].text == "08/04/2026"
    assert tab.rows[1].cells[cab.index("Data de empenho")].text == "01/02/2026"
