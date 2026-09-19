"""CONSOLIDADO — a carteira inteira do cliente, lado a lado (18/09/2026).

Pedido da assessoria Freitas: "preciso de uma lista do Reginaldo Lopes, onde ele
mandou $$ para nossos clientes". Nenhuma tela respondia: todas as operacionais
exigem UM município.

⚠️ REVERTE A DECISÃO DE 05/08/2026 SÓ NESTA FORMA. Naquele dia o dono tirou o
"Consolidado (todos)" do seletor de município: numa assessoria os municípios são
CLIENTES DIFERENTES, e somar as carteiras numa tela só cria a chance de ler o
número de um cliente achando que é de outro. O seletor continua sem "todos". O
que volta é uma ÁREA PRÓPRIA, para perguntas que atravessam clientes, e com uma
regra: todo número sai QUEBRADO POR MUNICÍPIO — total sozinho não existe aqui.

⚠️ ESCOPO: SEMPRE `services/bi.resolve_scope`, nunca `municipio_id=None` solto.
As funções de agregação reusadas (`aggregate_parlamentares`, `detalhe_core`) não
têm gate, e "sem escopo" nelas seria o tenant inteiro. Super-admin vê os
municípios ativos; os demais, os seus `user_municipios`.

⚠️ MESMA CONTA DAS OUTRAS TELAS: o ranking é `aggregate_parlamentares` e o
detalhe é `detalhe_core`, os mesmos da tela Parlamentares — a soma por município
aqui é a soma que cada município vê lá.
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession

from database import get_db
from models.user import User
from routers.parlamentares import aggregate_parlamentares, detalhe_core
from services.audit import registrar
from services.auth import ensure_tela, get_current_user
from services.bi import anos_list, resolve_scope
from services.cadastro_parlamentar import cadastros_por_nome
from services.registro_rotas import exige

router = APIRouter(prefix="/api/consolidado", tags=["consolidado"])

# Cada fonte do detalhe, com o campo de valor que o AGREGADO soma para ela. É a
# mesma escolha do `valor_total` de `detalhe_core` (voluntária = valor global,
# emenda estadual = valor da indicação, as demais = valor total): usar outro
# campo faria a soma por município divergir do número do cartão.
FONTES: tuple = (
    ("sigcon", "valor_total", "SIGCON (estadual)"),
    ("voluntarias", "valor_global", "Voluntárias (TransfereGov)"),
    ("emendas", "valor_indicacao", "Emendas estaduais"),
    ("plano_acao", "valor_total", "Emenda Pix (transferência especial)"),
    ("pac", "valor_total", "Novo PAC"),
    ("fns", "valor_total", "Saúde (FNS)"),
    ("emendas_federais", "valor_total", "Emendas federais (carteira CGU)"),
)


async def _escopo(db: AsyncSession, current: User) -> list[int]:
    ensure_tela(current, "consolidado")
    ids, _ = await resolve_scope(db, current, None)
    return ids


@router.get("/painel", dependencies=[exige("consolidado.ver")])
async def painel(
    db: AsyncSession = Depends(get_db),
    current: User = Depends(get_current_user),
):
    """A home do CONSOLIDADO: uma linha por município com regularidade,
    vencimentos, documentos, emendas sem pagamento e Radar — cada bloco com a
    conta da tela que já existe (ver `services/consolidado_painel.py`)."""
    from services.consolidado_painel import montar
    ids = await _escopo(db, current)
    return await montar(db, ids)


def por_municipio(det: dict) -> list[dict]:
    """Os lançamentos do detalhe agrupados por município. Função PURA.

    Cada município sai com o valor total, o número de lançamentos e a quebra por
    fonte (valor e contagem). Ordem: maior valor primeiro.
    """
    muns: dict = {}
    for fonte, campo, _rotulo in FONTES:
        for item in det.get(fonte) or []:
            mid = item.get("municipio_id")
            m = muns.setdefault(mid, {"municipio_id": mid,
                                      "municipio_nome": item.get("municipio_nome") or "—",
                                      "valor": 0.0, "lancamentos": 0, "por_fonte": {}})
            v = float(item.get(campo) or 0)
            m["valor"] += v
            m["lancamentos"] += 1
            f = m["por_fonte"].setdefault(fonte, {"valor": 0.0, "lancamentos": 0})
            f["valor"] += v
            f["lancamentos"] += 1
    return sorted(muns.values(), key=lambda m: (-m["valor"], m["municipio_nome"]))


def lancamentos(det: dict) -> list[dict]:
    """Todos os lançamentos do detalhe numa lista só, no formato da planilha. PURA."""
    saida = []
    for fonte, campo, rotulo in FONTES:
        for i in det.get(fonte) or []:
            saida.append({
                "fonte": rotulo,
                "municipio_id": i.get("municipio_id"),
                "municipio": i.get("municipio_nome") or "—",
                "numero": (i.get("numero") or i.get("numero_proposta") or i.get("codigo")
                           or i.get("nr_indicacao") or i.get("codigo_emenda") or ""),
                "ano": i.get("ano"),
                "objeto": (i.get("objeto") or i.get("beneficiario") or i.get("programa")
                           or i.get("beneficiario_nome") or ""),
                "situacao": (i.get("situacao") or i.get("status_indicacao") or ""),
                "valor": float(i.get(campo) or 0),
            })
    return sorted(saida, key=lambda x: (x["municipio"], -x["valor"]))


@router.get("/parlamentares", dependencies=[exige("consolidado.ver")])
async def parlamentares(
    q: Optional[str] = Query(None, description="Busca parcial no nome"),
    anos: Optional[list[int]] = Query(None),
    db: AsyncSession = Depends(get_db),
    current: User = Depends(get_current_user),
):
    """Quem destinou recurso à carteira: o ranking da tela Parlamentares, com o
    escopo de TODOS os municípios do usuário e só pessoas (sem Fundo Municipal)."""
    ids = await _escopo(db, current)
    dados = await aggregate_parlamentares(db, municipio_ids=ids, q=q, ano=anos_list(anos),
                                          tipo="parlamentar")
    cadastros = await cadastros_por_nome(db, {i["nome_display"] for i in dados["items"]})
    for i in dados["items"]:
        i["cadastro"] = cadastros.get(i["nome_display"])
        i["qt_municipios"] = len(i.get("municipios") or [])
    return {"municipios_na_carteira": len(ids), "items": dados["items"],
            "total": len(dados["items"])}


@router.get("/parlamentares/{nome:path}/exportar", dependencies=[exige("consolidado.exportar")])
async def exportar_parlamentar(
    request: Request,
    nome: str,
    formato: str = Query("xlsx", pattern="^(pdf|xlsx)$"),
    anos: Optional[list[int]] = Query(None),
    db: AsyncSession = Depends(get_db),
    current: User = Depends(get_current_user),
):
    """O relatório "onde o parlamentar mandou dinheiro na carteira", em arquivo:
    uma tabela por município e a lista de lançamentos."""
    ids = await _escopo(db, current)
    _anos = anos_list(anos)
    det = await detalhe_core(db, nome, ids, _anos)
    muns = por_municipio(det)
    lanc = lancamentos(det)
    periodo = ", ".join(map(str, _anos)) if _anos else "todos os anos"
    base = "".join(ch for ch in nome if ch.isalnum())[:30] or "parlamentar"
    arquivo = f"consolidado_{base}.{formato}"
    buf = (_xlsx(nome, periodo, len(ids), muns, lanc) if formato == "xlsx"
           else _pdf(nome, periodo, len(ids), muns, lanc))
    await registrar(
        db, action="export.consolidado_parlamentar", user=current, request=request,
        target_type="export", target_id="consolidado_parlamentar", alvo_nome=arquivo,
        details={"formato": formato, "registros": len(lanc), "arquivo": arquivo,
                 "filtros": {"parlamentar": nome, "anos": _anos,
                             "escopo": f"{len(ids)} municípios da carteira"}},
    )
    tipo = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            if formato == "xlsx" else "application/pdf")
    return StreamingResponse(buf, media_type=tipo,
                             headers={"Content-Disposition": f"attachment; filename={arquivo}"})


# ⚠️ DECLARADA DEPOIS DE `/exportar`: `{nome:path}` casa qualquer coisa, inclusive
# "fulano/exportar", e o FastAPI casa na ordem de registro.
@router.get("/parlamentares/{nome:path}", dependencies=[exige("consolidado.ver")])
async def parlamentar(
    nome: str,
    anos: Optional[list[int]] = Query(None),
    db: AsyncSession = Depends(get_db),
    current: User = Depends(get_current_user),
):
    """Onde este parlamentar mandou dinheiro na carteira: por município (valor e
    fontes) e os lançamentos, que são os mesmos da tela Parlamentares."""
    ids = await _escopo(db, current)
    det = await detalhe_core(db, nome, ids, anos_list(anos))
    return {"nome": nome, "municipios_na_carteira": len(ids),
            "por_municipio": por_municipio(det), "lancamentos": lancamentos(det),
            "valor_total": det["valor_total"], "total_lancamentos": det["total_geral"],
            "fontes": [{"chave": f, "rotulo": r} for f, _c, r in FONTES]}


# --------------------------------------------------------------- arquivos ---

def _br(v: float) -> str:
    return f"R$ {v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _xlsx(nome: str, periodo: str, n_carteira: int, muns: list[dict],
          lanc: list[dict]) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Por município"
    ws.append([f"{nome} — onde destinou recurso na carteira ({periodo})"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([f"{len(muns)} de {n_carteira} municípios da carteira · "
               f"gerado em {datetime.now():%d/%m/%Y %H:%M} · PACTHA"])
    ws.append([])
    rotulos = [r for _f, _c, r in FONTES]
    ws.append(["Município", "Valor total", "Lançamentos"] + rotulos)
    for c in ws[4]:
        c.font = Font(bold=True)
    for m in muns:
        ws.append([m["municipio_nome"], round(m["valor"], 2), m["lancamentos"]]
                  + [round(m["por_fonte"].get(f, {}).get("valor", 0.0), 2) for f, _c, _r in FONTES])
    for row in ws.iter_rows(min_row=5, min_col=2, max_col=2):
        for c in row:
            c.number_format = '"R$" #,##0.00'
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 18

    wl = wb.create_sheet("Lançamentos")
    wl.append(["Município", "Fonte", "Número", "Ano", "Situação", "Valor", "Objeto"])
    for c in wl[1]:
        c.font = Font(bold=True)
    for x in lanc:
        wl.append([x["municipio"], x["fonte"], x["numero"], x["ano"], x["situacao"],
                   round(x["valor"], 2), x["objeto"]])
    for row in wl.iter_rows(min_row=2, min_col=6, max_col=6):
        for c in row:
            c.number_format = '"R$" #,##0.00'
    for col, w in zip("ABCDEFG", (28, 34, 18, 8, 30, 16, 80)):
        wl.column_dimensions[col].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _pdf(nome: str, periodo: str, n_carteira: int, muns: list[dict],
         lanc: list[dict]) -> BytesIO:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    st = getSampleStyleSheet()
    titulo = ParagraphStyle("t", parent=st["Heading1"], fontSize=14,
                            textColor=colors.HexColor("#1e40af"), spaceAfter=2)
    sub = ParagraphStyle("s", parent=st["Normal"], fontSize=9,
                         textColor=colors.HexColor("#475569"), spaceAfter=8)
    cel = ParagraphStyle("c", parent=st["Normal"], fontSize=7, leading=8.5)

    def tabela(dados, larguras):
        t = Table(dados, repeatRows=1, colWidths=[w * mm for w in larguras], hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e40af")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
        ]))
        return t

    def p(txt, limite=300):
        s = str(txt or "")
        return Paragraph(escape(s[:limite] + ("…" if len(s) > limite else "")), cel)

    story = [
        Paragraph(escape(f"{nome} — onde destinou recurso na carteira"), titulo),
        Paragraph(escape(f"{periodo} · {len(muns)} de {n_carteira} municípios da carteira · "
                         f"{len(lanc)} lançamento(s)"), sub),
        tabela([["Município", "Valor", "Lançamentos", "Fontes"]]
               + [[p(m["municipio_nome"]), _br(m["valor"]), str(m["lancamentos"]),
                   p(" · ".join(f"{r}: {_br(m['por_fonte'][f]['valor'])}"
                                for f, _c, r in FONTES if f in m["por_fonte"]))]
                  for m in muns],
               [60, 32, 22, 160]),
        Spacer(1, 10),
        Paragraph("Lançamentos", ParagraphStyle("h", parent=st["Heading2"], fontSize=11)),
        tabela([["Município", "Fonte", "Número", "Ano", "Situação", "Valor", "Objeto"]]
               + [[p(x["municipio"]), p(x["fonte"]), p(x["numero"], 30), str(x["ano"] or ""),
                   p(x["situacao"], 60), _br(x["valor"]), p(x["objeto"])] for x in lanc],
               [34, 36, 22, 10, 34, 24, 117]),
        Spacer(1, 8),
        Paragraph(f"Gerado em {datetime.now():%d/%m/%Y %H:%M} | PACTHA", sub),
    ]
    buf = BytesIO()
    SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                      topMargin=10 * mm, bottomMargin=10 * mm).build(story)
    buf.seek(0)
    return buf
