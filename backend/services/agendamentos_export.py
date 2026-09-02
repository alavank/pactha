"""A relação de AGENDAMENTOS em arquivo — Excel e PDF do MESMO recorte.

O dono pediu "exportar a relação com todos os dados, conforme o filtro". Quem
decide QUAIS linhas entram é o router (`routers/agendamentos._filtros`, a mesma
função que a tela usa); este módulo só DESENHA o que recebe.

⚠️ ESTE MÓDULO NÃO FILTRA E NÃO CALCULA NADA. É a mesma disciplina do
`vigencias_export`: um recorte, dois renderizadores. Se cada formato montasse a
própria consulta, o Excel e o PDF divergiriam no primeiro ajuste do filtro — e a
divergência aparece como "o relatório veio com um agendamento a mais", sem nada
para culpar. Aqui os dois recebem a MESMA lista de dicionários.

⚠️ A DATA VAI NATIVA NO EXCEL, não como texto. Quem pede Excel vai ordenar e
filtrar por data; uma coluna de datas em texto ordena "01/12" antes de "02/01" e
o autofiltro não oferece o seletor de período. É o mesmo cuidado que o
`vigencias_export` toma com os valores — coluna em texto responde errado sem
avisar, que é o defeito mais silencioso de uma planilha.

⚠️ O RELATO VAI INTEIRO, e é ele que faz o arquivo existir. "Todos os dados"
inclui o texto da atividade: uma relação com título e data e sem o relato serve
para conferir agenda, não para prestar contas do que foi feito. No PDF ele é
quebrado em parágrafo; no Excel vai numa célula com quebra de linha.

⚠️ O ANEXO NÃO VAI NO ARQUIVO — vai o NOME dele. O binário sairia da plataforma
sem passar pela permissão `agendamentos.anexo_baixar`, que existe justamente
para separar "ver que há um anexo" de "receber o documento". Uma planilha com os
PDFs embutidos seria uma porta lateral para o cofre de documentos.
"""
from __future__ import annotations

import io
from datetime import date, datetime

_AZUL = "1E40AF"
_CINZA = "475569"

# As cores das três colunas do kanban, para a planilha e o PDF contarem a mesma
# história que a tela. Tons claros: a célula é fundo de texto, não etiqueta.
_COR_STATUS = {"a_fazer": "E8EDF5", "em_andamento": "FCF0DA", "realizado": "E4F1E8"}

_COLUNAS = [
    ("data", "Data", 12),
    ("municipio", "Município", 26),
    ("titulo", "Título", 38),
    ("status_rotulo", "Situação", 15),
    ("responsavel", "Responsável", 24),
    ("relato", "Relato da atividade", 60),
    ("anexos_nomes", "Anexos", 30),
    ("criado_por_nome", "Registrado por", 22),
]


def _dia(iso: str | None) -> date | None:
    """`YYYY-MM-DD` -> date. Devolve None em qualquer coisa ilegível.

    ⚠️ NÃO usa `datetime.fromisoformat` sobre um valor com hora: a coluna do
    banco é DATE justamente para não ter fuso, e reintroduzir um instante aqui
    traria de volta o problema de a data virar o dia anterior."""
    if not iso:
        return None
    try:
        return datetime.strptime(str(iso)[:10], "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def _linha(item: dict) -> dict:
    """Um agendamento no formato das colunas, sem nada calculado."""
    anexos = item.get("anexos") or []
    return {
        **item,
        "anexos_nomes": ", ".join(a.get("nome", "") for a in anexos) or "",
        "_data": _dia(item.get("data")),
    }


def gerar_xlsx(itens: list[dict]) -> bytes:
    """Uma aba, uma linha por agendamento, com autofiltro."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Agendamentos"

    azul = PatternFill("solid", fgColor=_AZUL)
    cab_fonte = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    normal = Font(name="Calibri", size=10)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    esq = Alignment(horizontal="left", vertical="top", wrap_text=True)
    fino = Side(style="thin", color="BFBFBF")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)

    for i, (_, rotulo, largura) in enumerate(_COLUNAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
        c = ws.cell(1, i, rotulo)
        c.fill = azul; c.font = cab_fonte; c.alignment = centro; c.border = borda
    ws.freeze_panes = "A2"
    # Autofiltro no cabeçalho: é o motivo de alguém pedir Excel em vez de PDF.
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_COLUNAS))}1"

    for r, item in enumerate((_linha(i) for i in itens), start=2):
        cor = _COR_STATUS.get(item.get("status") or "")
        fundo = PatternFill("solid", fgColor=cor) if cor else None
        for i, (chave, _, _) in enumerate(_COLUNAS, start=1):
            valor = item["_data"] if chave == "data" else item.get(chave)
            c = ws.cell(r, i, valor if valor is not None else "")
            c.font = normal
            c.border = borda
            c.alignment = centro if chave in ("data", "status_rotulo") else esq
            if chave == "data" and item["_data"]:
                c.number_format = "DD/MM/YYYY"
            if fundo and chave == "status_rotulo":
                c.fill = fundo

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def gerar_pdf(itens: list[dict], *, de: date | None = None,
              ate: date | None = None) -> bytes:
    """A relação em retrato, uma linha por agendamento.

    ⚠️ RETRATO, e não paisagem como os outros PDFs do repo. Os exports de
    convênios são tabelas de números — muitas colunas estreitas. Aqui a coluna
    que importa é o RELATO, um texto corrido: em paisagem ele vira uma faixa
    baixa e larga, difícil de ler, e o documento é justamente o que se anexa a
    um processo para contar o que foi feito.
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=12 * mm,
                            rightMargin=12 * mm, topMargin=12 * mm,
                            bottomMargin=12 * mm)
    estilos = getSampleStyleSheet()
    t_titulo = ParagraphStyle("T", parent=estilos["Heading1"], fontSize=14,
                              textColor=colors.HexColor(f"#{_AZUL}"), spaceAfter=4)
    t_sub = ParagraphStyle("S", parent=estilos["Normal"], fontSize=9,
                           textColor=colors.HexColor(f"#{_CINZA}"), spaceAfter=8)
    t_cel = ParagraphStyle("C", parent=estilos["Normal"], fontSize=8, leading=10)
    t_cab = ParagraphStyle("H", parent=estilos["Normal"], fontSize=8, leading=10,
                           textColor=colors.white)

    periodo = ""
    if de and ate:
        periodo = f"Período de {de:%d/%m/%Y} a {ate:%d/%m/%Y}. "
    elif de:
        periodo = f"A partir de {de:%d/%m/%Y}. "
    elif ate:
        periodo = f"Até {ate:%d/%m/%Y}. "
    # ⚠️ O SUBTÍTULO DIZ O RECORTE. Um relatório sem o filtro escrito nele vira
    # uma lista sem contexto assim que sai do navegador — e quem recebe não tem
    # como saber se está vendo o mês, o ano ou tudo.
    sub = (f"{periodo}{len(itens)} agendamento(s). "
           f"Gerado em {date.today():%d/%m/%Y}.")

    historia = [Paragraph("Relação de agendamentos", t_titulo),
                Paragraph(sub, t_sub), Spacer(1, 4)]

    cabecalho = [Paragraph(f"<b>{r}</b>", t_cab)
                 for r in ("Data", "Município", "Título / Relato",
                           "Situação", "Responsável")]
    linhas = [cabecalho]
    for item in (_linha(i) for i in itens):
        d = item["_data"]
        corpo = f"<b>{_esc(item.get('titulo'))}</b>"
        if item.get("relato"):
            corpo += f"<br/>{_esc(item['relato'])}"
        if item.get("anexos_nomes"):
            corpo += f"<br/><i>Anexos: {_esc(item['anexos_nomes'])}</i>"
        linhas.append([
            Paragraph(f"{d:%d/%m/%Y}" if d else "—", t_cel),
            Paragraph(_esc(item.get("municipio")), t_cel),
            Paragraph(corpo, t_cel),
            Paragraph(_esc(item.get("status_rotulo")), t_cel),
            Paragraph(_esc(item.get("responsavel")) or "—", t_cel),
        ])

    tabela = Table(linhas, colWidths=[20 * mm, 32 * mm, 78 * mm, 24 * mm, 32 * mm],
                   repeatRows=1)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_AZUL}")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    # A cor da situação é a mesma do kanban — o PDF conta a mesma história que a
    # tela de onde ele saiu.
    for i, item in enumerate(itens, start=1):
        cor = _COR_STATUS.get(item.get("status") or "")
        if cor:
            estilo.append(("BACKGROUND", (3, i), (3, i), colors.HexColor(f"#{cor}")))
    tabela.setStyle(TableStyle(estilo))
    historia.append(tabela)

    if not itens:
        historia.append(Spacer(1, 6))
        historia.append(Paragraph(
            "Nenhum agendamento no recorte selecionado.", t_sub))

    doc.build(historia)
    return buf.getvalue()


def _esc(v) -> str:
    """Texto do usuário dentro de `Paragraph` do reportlab.

    ⚠️ O reportlab interpreta um subconjunto de HTML no `Paragraph`. Um relato
    que contenha `<` ou `&` — e relato é campo livre — quebra a construção do
    PDF com um erro de parser, ou some da página. Escapar é obrigatório, e a
    quebra de linha do texto vira `<br/>` porque o `Paragraph` ignora `\\n`.
    """
    import html
    if v is None:
        return ""
    return html.escape(str(v)).replace("\n", "<br/>")
