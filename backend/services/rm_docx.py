"""Renderizador WORD (.docx) do RM — o gemeo de `services/rm_pdf.gerar_pdf`.

⚠️ ESTE ARQUIVO NAO DECIDE CONTEUDO. Ele consome `services.rm_pdf.roteiro_rm`,
exatamente a mesma funcao que o PDF consome, entao os dois documentos trazem as
MESMAS partes, secoes, orgaos, itens, campos e caixas de destaque, na mesma
ordem. Campo novo ou caixa nova entram no roteiro e aparecem aqui de graca.

O que e responsabilidade DESTE arquivo (e so isto):
  - traduzir os 10 estilos de `rm_pdf._styles()` para formatacao do Word;
  - desenhar as caixas ambar/cinza com `w:pBdr` + `w:shd` (ver `_moldura`);
  - traduzir o mini-HTML das caixas (<b>, <i>, <br/>) para runs do Word;
  - cabecalho com logo e rodape com numero de pagina, que no PDF vivem fora da
    story (`rm_pdf._on_page`) e por isso nao passam pelo roteiro.

⚠️ A PAGINACAO NAO BATE PAGINA A PAGINA COM O PDF, e isso nao tem conserto: quem
quebra linha aqui e o motor do Word, nao o ReportLab. Mesmo conteudo, mesma
ordem, mesmas medidas — numero de paginas possivelmente diferente. Embutir o PDF
dentro do .docx daria paginacao identica e tiraria justamente o que se pediu, que
e um documento EDITAVEL.
"""
from __future__ import annotations
import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from services.rm_pdf import (
    _CAB_DIST_CM, _GLIFO_CAMPO, _GLIFO_GRUPO, _LOGO_ALT_CM, _MARGENS_CM,
    _ROD_DIST_CM, _logo_path, roteiro_rm,
)

# ⚠️ ORDEM DO SCHEMA ECMA-376 — A ARMADILHA DESTE ARQUIVO.
#
# `w:pPr` NAO e um saco de filhos: a sequencia de CT_PPr e FIXA. `w:pBdr` vem
# ANTES de `w:shd`, e os DOIS vem antes de `w:spacing`, `w:ind` e `w:jc` — que sao
# justamente o que `paragraph_format` escreve quando ajustamos espacamento e
# recuo. Um `pPr.append(borda)` cego encosta a borda DEPOIS do espacamento; o
# Word entao ou abre a caixa de dialogo "conteudo ilegivel" ou DESCARTA o
# elemento em silencio — e o silencio e o caso pior, porque a caixa ambar da
# clausula suspensiva sumiria do documento oficial sem ninguem perceber.
#
# python-docx NAO expoe `pBdr`/`shd` em CT_PPr (nao ha `get_or_add_pBdr`), entao a
# insercao e manual via `insert_element_before`, com a lista de SUCESSORES abaixo.
# As duas listas sao recortes literais de `docx.oxml.text.parfmt.CT_PPr._tag_seq`
# (que a propria classe apaga com `del` no fim, e por isso nao da para ler em
# tempo de execucao).
_SUCESSORES_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)
_SUCESSORES_SHD = _SUCESSORES_PBDR[1:]      # tudo o que vem depois de w:shd
_SUCESSORES_TABS = _SUCESSORES_PBDR[2:]     # tudo o que vem depois de w:tabs

# FONTE: o PDF usa Helvetica. O equivalente metrico no Word e Arial — as mesmas
# larguras de glifo, entao a quebra de linha cai praticamente no mesmo lugar.
_FONTE = "Arial"


def _hex(cor: str) -> RGBColor:
    return RGBColor.from_string(cor.replace("#", "").upper())


# ---------------------------------------------------------------------------
# ESTILOS — a traducao 1:1 de `rm_pdf._styles()`.
#
# ⚠️ ENTRELINHA 1,5 EM TODO O DOCUMENTO. No PDF isso e `leading = 1,5 x fontSize`
# (ver a docstring de `rm_pdf._styles`); aqui e `line_spacing = 1.5`, que e a
# mesma regra escrita do jeito do Word — e, de fato, e de um .docx que a medida
# saiu (`w:spacing w:line="360"`, 360/240 = 1,5).
#
# `rodape` e o unico estilo que existe nos dois lados e so AQUI e usado de
# verdade: no PDF o rodape e desenhado com chamadas cruas de canvas
# (`rm_pdf._on_page`) e `s["rodape"]` fica inerte.
#
# ⚠️ `rodape` NAO leva `alin`: quem posiciona o texto no rodape sao as TABULACOES
# (ver `_montar_secao`), e o paragrafo tem de ficar alinhado a esquerda para elas
# valerem. Pedir "centro" aqui seria uma instrucao que a linha seguinte desfaz —
# comentario mentiroso esperando para confundir o proximo leitor.
# ---------------------------------------------------------------------------
_ESTILOS = {
    "titulo_principal": dict(tam=13, negrito=True, alin="centro", depois=4),
    "data_local":       dict(tam=11, alin="direita", depois=14),
    "parte_titulo":     dict(tam=12, negrito=True, alin="centro", antes=14, depois=10,
                             cor="#1e40af"),
    "secao_titulo":     dict(tam=11, negrito=True, alin="centro", antes=10, depois=8,
                             cor="#111827"),
    "grupo_titulo":     dict(tam=10.5, negrito=True, antes=10, depois=4, esq=4),
    "item_id":          dict(tam=10, negrito=True, antes=14, depois=2, esq=14),
    "item_campo":       dict(tam=9.5, alin="justificado", depois=1, esq=28),
    "rodape":           dict(tam=8, cor="#475569"),
    "clausula":         dict(tam=9.5, esq=28, dir=10, antes=3, depois=3,
                             cor="#7c2d12", fundo="#FEF3C7", borda="#D97706"),
    "informativo":      dict(tam=9.5, esq=28, dir=10, antes=3, depois=3,
                             cor="#334155", fundo="#F1F5F9", borda="#CBD5E1"),
}

_ALINHAMENTO = {
    "centro": WD_ALIGN_PARAGRAPH.CENTER,
    "direita": WD_ALIGN_PARAGRAPH.RIGHT,
    "justificado": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _moldura(p, fundo: str, borda: str) -> None:
    """Caixa de destaque: fundo (`w:shd`) + borda nos 4 lados (`w:pBdr`).

    A borda sai com `w:sz` em OITAVOS DE PONTO (o PDF usa borderWidth=1pt, logo
    sz=8) e `w:space` em PONTOS (o PDF usa borderPadding=5, logo space=5).

    ⚠️ A ordem de insercao e o motivo de esta funcao existir — ver o comentario
    de `_SUCESSORES_PBDR` no topo do arquivo."""
    pPr = p._p.get_or_add_pPr()

    pbdr = OxmlElement("w:pBdr")
    for lado in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")        # 8 oitavos = 1pt (borderWidth do PDF)
        el.set(qn("w:space"), "5")     # borderPadding do PDF
        el.set(qn("w:color"), borda.replace("#", "").upper())
        pbdr.append(el)
    pPr.insert_element_before(pbdr, *_SUCESSORES_PBDR)

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fundo.replace("#", "").upper())
    pPr.insert_element_before(shd, *_SUCESSORES_SHD)


def _limpar_tabulacoes(p) -> None:
    """Apaga as tabulacoes HERDADAS do estilo `Footer` do template.

    ⚠️ `w:tabs` E CUMULATIVO NA HERANCA DE ESTILO. O estilo `Footer` do template
    padrao do python-docx traz duas paradas pensadas para papel CARTA (4680 e
    9360 twips = 8,25 cm e 16,5 cm). Somando-se as nossas, que sao calculadas
    para A4, o Word usa a PRIMEIRA parada de cada tabulacao encontrada: o
    endereco sairia ~0,5 cm fora do centro e o numero da pagina ~1 cm dentro da
    margem — em toda pagina de todo tenant, sem erro nenhum.

    `w:val="clear"` remove a parada herdada naquela posicao exata."""
    pPr = p._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    for pos in ("4680", "9360"):
        t = OxmlElement("w:tab")
        t.set(qn("w:val"), "clear")
        t.set(qn("w:pos"), pos)
        tabs.append(t)
    pPr.insert_element_before(tabs, *_SUCESSORES_TABS)


def _paragrafo(doc, estilo: str):
    """Paragrafo novo ja com o estilo aplicado. Devolve (paragrafo, spec)."""
    p = doc.add_paragraph()
    return _formatar(p, estilo), _ESTILOS[estilo]


def _formatar(p, estilo: str):
    spec = _ESTILOS[estilo]
    pf = p.paragraph_format
    pf.line_spacing = 1.5          # a regra global — ver bloco de comentario acima
    pf.space_before = Pt(spec.get("antes", 0))
    pf.space_after = Pt(spec.get("depois", 0))
    if spec.get("esq"):
        pf.left_indent = Pt(spec["esq"])
    if spec.get("dir"):
        pf.right_indent = Pt(spec["dir"])
    if spec.get("alin"):
        p.alignment = _ALINHAMENTO[spec["alin"]]
    # ⚠️ A moldura DEPOIS do paragraph_format de proposito: `insert_element_before`
    # so acha o lugar certo se os sucessores (w:spacing, w:ind, w:jc) ja estiverem
    # no XML. Na ordem inversa tambem funciona, mas assim o teste de ordem cobre o
    # caso ruim de verdade.
    if spec.get("fundo"):
        _moldura(p, spec["fundo"], spec["borda"])
    return p


def _espacador(doc, pontos: float):
    """Equivalente do `Spacer` do ReportLab: um paragrafo vazio de N pontos.

    ⚠️ NAO E ENFEITE, E CORRECAO DE UM DEFEITO DE VERDADE. Alem de reproduzir o
    `Spacer(1, 2)` que o PDF poe antes de cada caixa, ele impede que o Word FUNDA
    caixas vizinhas: paragrafos adjacentes com `w:pBdr` IDENTICO sao desenhados
    pelo Word como UMA moldura so. Sem este paragrafo, um item com evento +
    alteracao + desembolso + obra (quatro caixas cinza seguidas) sairia como um
    unico bloco cinza — conteudo certo, leitura errada. A caixa ambar nao funde
    com a cinza porque as bordas diferem, mas duas cinzas fundem."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    r = p.add_run("")
    r.font.name = _FONTE
    r.font.size = Pt(pontos)
    return p


def _run(p, texto: str, estilo: str, negrito=False, italico=False):
    spec = _ESTILOS[estilo]
    r = p.add_run(texto)
    r.font.name = _FONTE
    r.font.size = Pt(spec["tam"])
    r.bold = bool(spec.get("negrito")) or negrito
    r.italic = italico
    if spec.get("cor"):
        r.font.color.rgb = _hex(spec["cor"])
    return r


def _uma_linha(v) -> str:
    """Colapsa espaco, TAB e quebra de linha num espaco so.

    ⚠️ NAO E COSMETICO. O ReportLab COLAPSA `\\n` e `\\t` em espaco ao montar o
    Paragraph; o python-docx faz o oposto — converte em `<w:br/>` e `<w:tab/>`.
    O mesmo valor (o `objeto` de um convenio copiado do portal costuma trazer
    quebras) sairia numa linha no PDF e em tres no Word, com tabulacao no meio."""
    return " ".join(str(v).split())


# --------------------------------------------------------------- MINI-HTML ----
# As seis funcoes `_*_destaque` de rm_pdf devolvem o mini-HTML do Paragraph do
# ReportLab: <b>, <i> e <br/> como separador de linha, com &/</> ja escapados.
# O Word imprimiria isso LITERAL ("<b>Motivo:</b>"), entao o markup e traduzido
# para runs aqui. O dialeto e fechado e pequeno — nao ha <font>, <para>, <link>
# nem atributo nenhum no arquivo inteiro (conferido tag a tag em rm_pdf).
_TAGS = re.compile(r"(<b>|</b>|<i>|</i>)")


def _desescapar(s: str) -> str:
    """Inverso exato de `rm_pdf._escape`.

    ⚠️ `&amp;` POR ULTIMO. Na ordem inversa, o texto original "&lt;" (que virou
    "&amp;lt;") sairia como "<" — o escape seria desfeito duas vezes."""
    return s.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")


def _runs_do_markup(p, markup: str, estilo: str) -> None:
    negrito = italico = False
    for i, linha in enumerate(markup.split("<br/>")):
        if i:
            # quebra de linha DENTRO do mesmo paragrafo — a caixa tem de ser UM
            # paragrafo so, senao o Word desenha uma moldura por linha.
            _run(p, "", estilo).add_break()
        for pedaco in _TAGS.split(linha):
            if pedaco == "<b>":
                negrito = True
            elif pedaco == "</b>":
                negrito = False
            elif pedaco == "<i>":
                italico = True
            elif pedaco == "</i>":
                italico = False
            elif pedaco:
                # `_uma_linha` pela mesma razao do campo: `\n`/`\t` que sobraram
                # no texto da caixa viram `<w:br/>`/`<w:tab/>` no Word e espaco no
                # PDF. A quebra INTENCIONAL e `<br/>`, tratada acima.
                _run(p, _uma_linha(_desescapar(pedaco)), estilo,
                     negrito=negrito, italico=italico)


# ------------------------------------------------- CABECALHO E RODAPE ---------
def _campo_pagina(p, estilo: str) -> None:
    """Numero de pagina como CAMPO do Word (`PAGE`), nao como texto fixo."""
    r = _run(p, "", estilo)
    ini = OxmlElement("w:fldChar"); ini.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fim = OxmlElement("w:fldChar"); fim.set(qn("w:fldCharType"), "end")
    r._r.append(ini); r._r.append(instr); r._r.append(fim)


def _montar_secao(doc, rodape_txt: str) -> None:
    """Pagina, margens, cabecalho com logo e rodape — o que no PDF esta em
    `rm_pdf.gerar_pdf` (SimpleDocTemplate) e em `rm_pdf._on_page`.

    ⚠️ O rodape e o numero de pagina NAO passam pelo roteiro porque no PDF eles
    nao estao na story: sao canvas cru. Um renderizador que so traduzisse a story
    sairia com todas as paginas sem rodape, e sem erro nenhum."""
    sec = doc.sections[0]
    # A4 — o template padrao do python-docx e CARTA (Letter). Sem estas duas
    # linhas o Word abriria o RM em 21,59 x 27,94 cm.
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    # As MESMAS medidas do PDF, lidas da mesma constante (rm_pdf._MARGENS_CM):
    # topo, base, esquerda, direita. Cravar os numeros aqui faria "aumentar a
    # margem" ser corrigido num formato so.
    sec.top_margin = Cm(_MARGENS_CM[0])
    sec.bottom_margin = Cm(_MARGENS_CM[1])
    sec.left_margin = Cm(_MARGENS_CM[2])
    sec.right_margin = Cm(_MARGENS_CM[3])
    sec.header_distance = Cm(_CAB_DIST_CM)
    sec.footer_distance = Cm(_ROD_DIST_CM)
    # `different_first_page_header_footer` fica FALSO (o padrao): cabecalho e
    # rodape valem desde a 1a pagina, como no PDF. O .docx da referencia nao
    # marca `titlePg` — e a mesma razao citada em `rm_pdf._on_page`.
    sec.different_first_page_header_footer = False

    cab = sec.header.paragraphs[0]
    cab.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _logo = _logo_path()
    if _logo:
        try:
            # A mesma altura do PDF (`rm_pdf._LOGO_ALT_CM`). Aqui so a ALTURA e
            # fixada (a largura acompanha a proporcao); o PDF ainda limita a
            # largura ao dobro. Nenhum dos quatro brasoes/logos passa de 2:1,
            # entao na pratica os dois cabecalhos saem do mesmo tamanho.
            cab.add_run().add_picture(_logo, height=Cm(_LOGO_ALT_CM))
        except Exception:
            pass          # logo ilegivel nunca derruba a emissao (igual ao PDF)

    rod = sec.footer.paragraphs[0]
    _formatar(rod, "rodape")
    rod.alignment = WD_ALIGN_PARAGRAPH.LEFT     # quem posiciona sao as tabulacoes
    rod.paragraph_format.line_spacing = 1.0
    # Rodape CENTRALIZADO e numero de pagina A DIREITA na MESMA linha, como no
    # PDF (`drawCentredString` + `drawRightString`). No Word isso e uma tabulacao
    # centrada no meio da area util e outra alinhada a direita na borda dela.
    _limpar_tabulacoes(rod)          # ⚠️ ANTES de por as nossas — ver a funcao
    util = 21.0 - _MARGENS_CM[2] - _MARGENS_CM[3]      # area util em cm
    tabs = rod.paragraph_format.tab_stops
    tabs.add_tab_stop(Cm(util / 2), WD_TAB_ALIGNMENT.CENTER)
    tabs.add_tab_stop(Cm(util), WD_TAB_ALIGNMENT.RIGHT)
    _run(rod, "\t" + (rodape_txt or "") + "\t", "rodape")
    _campo_pagina(rod, "rodape")


# ------------------------------------------------------------- RENDERIZADOR ---
def gerar_docx_rm(meta: dict, conteudo: dict, municipio_nome: str) -> bytes:
    """Gera bytes do .docx do RM.

    MESMA ASSINATURA de `rm_pdf.gerar_pdf` de proposito: e o que permite ao
    router e ao recorte do Resumido trocarem de formato mudando uma linha.

    Args:
        meta: {data_referencia, cidade_emissao, titulo (opt), rodape, escopo (opt)}
        conteudo: {partes: [...]}  — o mesmo JSONB que o PDF recebe
        municipio_nome: para o titulo
    """
    doc = Document()
    base = doc.styles["Normal"]
    base.font.name = _FONTE
    base.font.size = Pt(9.5)
    _montar_secao(doc, meta.get("rodape", ""))

    quebrar_proxima_parte = False
    # ⚠️ CONTA ELEMENTOS, NAO CAMPOS. O PDF faz `KeepTogether(bloco[:3])`, e
    # `bloco` guarda campos E os Spacer das caixas — logo sao os DOIS primeiros
    # elementos depois do identificador que ficam presos a ele.
    #
    # A aritmetica do Word: `keep_with_next` em P prende P ao paragrafo SEGUINTE.
    # Para manter {id, e0, e1} juntos bastam id.kwn e e0.kwn — e1.kwn puxaria um
    # QUARTO paragrafo, o que o PDF nao faz. Por isso o corte e `< 1`.
    elems_apos_id = 0

    for evento, dado in roteiro_rm(meta, conteudo, municipio_nome):
        if evento == "titulo":
            p, _ = _paragrafo(doc, "titulo_principal")
            _run(p, dado, "titulo_principal")
        elif evento == "local_data":
            p, _ = _paragrafo(doc, "data_local")
            _run(p, dado, "data_local")
        elif evento == "consultas":
            # "Consultas incluídas: ..." — só existe em RM filtrado. Mesmo estilo
            # da linha de cima, como no PDF.
            p, _ = _paragrafo(doc, "data_local")
            _run(p, dado, "data_local")
        elif evento == "quebra":
            # No Word a quebra vai NO PARAGRAFO da parte (`page_break_before`), e
            # nao num paragrafo vazio proprio: um paragrafo so para quebrar deixa
            # uma linha em branco no topo de toda parte a partir da 2a.
            quebrar_proxima_parte = True
        elif evento == "parte":
            p, _ = _paragrafo(doc, "parte_titulo")
            if quebrar_proxima_parte:
                p.paragraph_format.page_break_before = True
                quebrar_proxima_parte = False
            _run(p, dado, "parte_titulo")
        elif evento == "secao":
            p, _ = _paragrafo(doc, "secao_titulo")
            _run(p, dado, "secao_titulo")
        elif evento == "grupo":
            p, _ = _paragrafo(doc, "grupo_titulo")
            _run(p, f"{_GLIFO_GRUPO} {dado}", "grupo_titulo")
        elif evento == "item":
            p, _ = _paragrafo(doc, "item_id")
            # Equivalente do `KeepTogether(bloco[:3])` do PDF: o identificador
            # nunca fica orfao do proprio conteudo no fim da pagina. O resto do
            # item PODE quebrar entre paginas — que e exatamente o comportamento
            # que o PDF passou a ter (ver o comentario longo em `gerar_pdf`).
            p.paragraph_format.keep_with_next = True
            _run(p, dado, "item_id")
            elems_apos_id = 0
        elif evento == "campo":
            label, val = dado
            p, _ = _paragrafo(doc, "item_campo")
            if elems_apos_id < 1:
                p.paragraph_format.keep_with_next = True
            elems_apos_id += 1
            _run(p, f"{_GLIFO_CAMPO} ", "item_campo")
            _run(p, f"{label}: ", "item_campo", negrito=True)
            _run(p, _uma_linha(val), "item_campo")
        elif evento == "caixa":
            markup, estilo = dado
            esp = _espacador(doc, 2)   # o Spacer(1, 2) do PDF — e o anti-fusao
            if elems_apos_id < 1:
                esp.paragraph_format.keep_with_next = True
            elems_apos_id += 1
            p, _ = _paragrafo(doc, estilo)
            _runs_do_markup(p, markup, estilo)
        elif evento == "fim_item":
            _espacador(doc, 4)        # o Spacer(1, 4) que fecha o item no PDF
        elif evento == "vazio":
            p, _ = _paragrafo(doc, "item_campo")
            _run(p, dado, "item_campo", italico=True)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
