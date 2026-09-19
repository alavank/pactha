"""RELATÓRIOS DO CONSOLIDADO — PR 3 da série (19/09/2026).

Três relatórios da carteira, todos quebrados por município:
  - RECURSOS POR MUNICÍPIO: estaduais, voluntárias federais e emendas federais,
    lado a lado, no período escolhido;
  - MATRIZ PARLAMENTAR × MUNICÍPIO: quem mandou quanto para cada cliente;
  - a PLANILHA DO PAINEL da carteira.

⚠️ RECURSOS NÃO TEM COLUNA DE TOTAL, de propósito. As fontes se sobrepõem: a
voluntária que nasceu de emenda está nas voluntárias E nas emendas federais
(`emendas_unificadas` casa as duas pelo código). Somar as colunas contaria o
mesmo dinheiro duas vezes — e o total de um cliente seria o número errado.

⚠️ MESMA CONTA DAS TELAS: estaduais e voluntárias saem de `services/bi.bi_kpis`
(o Painel de cada município), emendas de `_fontes_federais` + `emendas_unificadas`
(a aba Federais), a matriz de `aggregate_parlamentares` (a tela Parlamentares),
que desde este PR devolve `por_municipio` numa varredura só.
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Optional

from sqlalchemy.ext.asyncio import AsyncSession

MOEDA = '"R$" #,##0.00'


async def recursos_por_municipio(db: AsyncSession, ids: list[int],
                                 anos: Optional[list[int]]) -> list[dict]:
    from routers.emendas_parlamentares import _fontes_federais
    from services.bi import bi_kpis
    from services.consolidado_painel import _municipios
    from services.emendas_unificadas import filtrar, totais, unificar_federais

    linhas = []
    for m in await _municipios(db, ids):
        mid = m["municipio_id"]
        k = await bi_kpis(db, [mid], anos)
        f = await _fontes_federais(db, mid)
        todas = unificar_federais((f["carteira"] or {}).get("items") or [], f["te"],
                                  f["parcerias"], f["indicadas"], f["voluntarias"])
        t = totais(filtrar(todas, anos=anos))
        linhas.append({
            "municipio_id": mid, "nome": m["nome"], "uf": m["uf"],
            "estaduais": {"n": k["total_convenios_estadual"], "valor": k["valor_total_estadual"]},
            "voluntarias": {"n": k["total_voluntarias"], "valor": k["valor_total_federal"]},
            "emendas": {"n": t["emendas"], "valor": t["valor_prefeitura"],
                        "sem_pagamento": t["parado_n"],
                        "estado": (f["carteira"] or {}).get("estado")},
        })
    return linhas


async def matriz_parlamentares(db: AsyncSession, ids: list[int], anos: Optional[list[int]],
                               limite: int = 200) -> dict:
    """Parlamentares (linhas) × municípios (colunas), numa agregação só."""
    from routers.parlamentares import aggregate_parlamentares
    from services.consolidado_painel import _municipios

    muns = await _municipios(db, ids)
    agg = await aggregate_parlamentares(db, municipio_ids=ids, ano=anos, tipo="parlamentar")
    itens = sorted(agg["items"], key=lambda i: -i["valor_total"])[:limite]
    return {"municipios": [m["nome"] for m in muns],
            "parlamentares": [{"nome": i["nome_display"], "total": i["valor_total"],
                               "por_municipio": i.get("por_municipio") or {}} for i in itens],
            "truncado": len(agg["items"]) > limite, "total_parlamentares": len(agg["items"])}


# --------------------------------------------------------------- planilhas ---

def _wb(titulo: str, subtitulo: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.append([titulo])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([f"{subtitulo} · gerado em {datetime.now():%d/%m/%Y %H:%M} · PACTHA"])
    ws.append([])
    return wb, ws


def _cabecalho(ws, colunas: list[str]) -> None:
    from openpyxl.styles import Font
    ws.append(colunas)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)


def _moeda(ws, colunas: list[int], desde: int) -> None:
    for col in colunas:
        for row in ws.iter_rows(min_row=desde, min_col=col, max_col=col):
            for c in row:
                c.number_format = MOEDA


def _salvar(wb) -> BytesIO:
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def xlsx_recursos(linhas: list[dict], periodo: str) -> BytesIO:
    wb, ws = _wb("Recursos por município da carteira",
                 f"{periodo} · {len(linhas)} municípios · sem coluna de total: as fontes se "
                 f"sobrepõem (voluntária que nasceu de emenda aparece nas duas)")
    ws.title = "Recursos por município"
    _cabecalho(ws, ["Município", "UF", "Convênios estaduais", "Valor estadual",
                    "Voluntárias federais", "Valor voluntárias", "Emendas federais",
                    "Valor das emendas (prefeitura)", "Emendas sem pagamento"])
    ini = ws.max_row + 1
    for l in linhas:
        ws.append([l["nome"], l["uf"], l["estaduais"]["n"], l["estaduais"]["valor"],
                   l["voluntarias"]["n"], l["voluntarias"]["valor"], l["emendas"]["n"],
                   l["emendas"]["valor"], l["emendas"]["sem_pagamento"]])
    _moeda(ws, [4, 6, 8], ini)
    for col, w in zip("ABCDEFGHI", (30, 5, 12, 18, 12, 18, 12, 22, 12)):
        ws.column_dimensions[col].width = w
    return _salvar(wb)


def xlsx_matriz(m: dict, periodo: str) -> BytesIO:
    from openpyxl.utils import get_column_letter
    wb, ws = _wb("Parlamentares × municípios da carteira",
                 f"{periodo} · {len(m['parlamentares'])} parlamentar(es)"
                 + (f" de {m['total_parlamentares']} (os de maior valor)" if m["truncado"] else ""))
    ws.title = "Matriz"
    muns = m["municipios"]
    _cabecalho(ws, ["Parlamentar", "Total na carteira"] + muns)
    ini = ws.max_row + 1
    for p in m["parlamentares"]:
        ws.append([p["nome"], p["total"]] + [p["por_municipio"].get(n) or None for n in muns])
    _moeda(ws, list(range(2, len(muns) + 3)), ini)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    for i in range(3, len(muns) + 3):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.freeze_panes = ws.cell(row=ini, column=3)
    return _salvar(wb)


def xlsx_painel(p: dict) -> BytesIO:
    """O painel da carteira em planilha: uma aba por bloco."""
    wb, ws = _wb("Painel da carteira", f"{p['municipios_na_carteira']} municípios")
    ws.title = "Município a município"
    _cabecalho(ws, ["Município", "UF", "CAUC", "Pendências CAUC", "Estadual",
                    "Convênios vencendo (90d)", "Próximo em (dias)", "Documentos vencendo (30d)",
                    "Emendas sem pagamento", "Radar: abertos", "Radar: nomeado",
                    "Radar: indicado"])
    for l in p["municipios"]:
        c, e = l.get("cauc"), l.get("estadual") or {}
        est = ("sem fonte no estado" if e.get("cobertura") == "sem_fonte"
               else "sem coleta" if e.get("regular") is None
               else "em dia" if e.get("regular") else "irregular")
        r, em = l.get("radar") or {}, l.get("emendas") or {}
        ws.append([l["nome"], l["uf"],
                   "sem coleta" if not c or c.get("regular") is None
                   else ("em dia" if c["regular"] else "irregular"),
                   (c or {}).get("pendencias"), est,
                   l["vencimentos"]["n"], l["vencimentos"]["proximo_dias"],
                   l["documentos"]["n"], em.get("sem_pagamento"),
                   r.get("abertos"), r.get("nomeado"), r.get("indicado")])
    ws.column_dimensions["A"].width = 30

    wv = wb.create_sheet("Vencimentos (90 dias)")
    _cabecalho(wv, ["Município", "Esfera", "Número", "Objeto", "Órgão", "Vence em", "Dias"])
    for v in p["vencimentos"]:
        wv.append([v["municipio"], v["esfera"], v["numero"], v["objeto"], v["orgao"],
                   v["fim"], v["dias"]])
    wd = wb.create_sheet("Documentos (30 dias)")
    _cabecalho(wd, ["Município", "Entidade", "Cadastro", "Documento", "Validade", "Dias"])
    for d in p["documentos"]:
        wd.append([d["municipio"], d.get("entidade"), d["esfera"], d["label"],
                   d["validade"], d["dias_restantes"]])
    wr = wb.create_sheet("Radar")
    _cabecalho(wr, ["Município", "Programa", "Município nomeado", "Fecha em (dias)",
                    "Emenda indicada por", "Valor indicado"])
    for x in p["radar"]:
        ind = x["indicacoes"] or [{}]
        for i in ind:
            wr.append([x["municipio"], x["programa"], "sim" if x["beneficiario"] else "",
                       x["dias"], i.get("parlamentar"), i.get("valor")])
    _moeda(wr, [6], 2)
    return _salvar(wb)
