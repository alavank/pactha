"""Render DOCX de um documento (módulo Geração de Documentos) a partir do
schema + dados preenchidos. Genérico: itera o schema, então serve p/ qualquer
tipo cadastrado em documentos_schema.py."""
from __future__ import annotations
import io

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _fmt(campo: dict, valor) -> str:
    if valor is None or valor == "":
        return "—"
    if campo.get("tipo") == "currency":
        try:
            v = float(str(valor).replace("R$", "").replace(".", "").replace(",", ".").strip())
            return f"R$ {v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        except (ValueError, TypeError):
            return str(valor)
    return str(valor)


def _add_campo(doc, label: str, valor: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{label}: ")
    r.bold = True
    # valores multi-linha: preserva quebras
    linhas = (valor or "—").split("\n")
    p.add_run(linhas[0])
    for ln in linhas[1:]:
        p.add_run().add_break()
        p.add_run(ln)


def _add_secao_titulo(doc, titulo: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(titulo.upper())
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)


def gerar_docx(schema: dict, dados: dict, municipio_nome: str = "", municipio_uf: str = "") -> bytes:
    dados = dados or {}
    doc = Document()

    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(11)

    # Título
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt = t.add_run(schema.get("titulo", "Documento").upper())
    rt.bold = True
    rt.font.size = Pt(16)
    if municipio_nome:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rs = sub.add_run(f"{municipio_nome}/{municipio_uf}".strip("/"))
        rs.font.size = Pt(11)
        rs.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    doc.add_paragraph()

    for secao in schema.get("secoes", []):
        _add_secao_titulo(doc, secao["titulo"])
        if secao.get("tipo") == "lista":
            itens = dados.get(secao.get("key")) or []
            if not itens:
                ip = doc.add_paragraph()
                ip.add_run("Nenhum item cadastrado.").italic = True
                continue
            for idx, item in enumerate(itens, 1):
                hp = doc.add_paragraph()
                hp.paragraph_format.space_before = Pt(6)
                hr = hp.add_run(f"{secao.get('item_label', 'Item')} {idx}")
                hr.bold = True
                hr.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                for campo in secao["campos"]:
                    _add_campo(doc, campo["label"], _fmt(campo, (item or {}).get(campo["key"])))
        else:
            for campo in secao["campos"]:
                _add_campo(doc, campo["label"], _fmt(campo, dados.get(campo["key"])))

    # Assinaturas
    if schema.get("rodape_assinatura"):
        doc.add_paragraph()
        nota = doc.add_paragraph()
        rn = nota.add_run(schema["rodape_assinatura"])
        rn.italic = True
        rn.font.size = Pt(9)
        rn.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        doc.add_paragraph()
        for label in ("Responsável pelo convenente", "Responsável pela sustentabilidade do objeto"):
            doc.add_paragraph()
            ln = doc.add_paragraph()
            ln.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ln.add_run("_" * 45)
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.add_run(label).font.size = Pt(10)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
