"""Ingestao do Acordo FES — divida do Fundo Estadual de Saude de MG (SES-MG)
com os credores da saude. Fonte: espelho Excel publico do Painel do Acordo FES
(saude.mg.gov.br/acordofes -> "Valores-acordo.xlsx").

Agrega por credor (CNPJ + razao social) e casa o "FUNDO MUNICIPAL DE SAUDE DE
<municipio>" aos nossos municipios (por nome normalizado, exato). Guarda todos
os credores (~1.4k) p/ busca livre por CNPJ/razao; tagueia os nossos.

Uso: DATABASE_URL_SYNC=... python ingestion/acordofes_ingest.py
"""
from __future__ import annotations
import os
import io
import re
import unicodedata
import logging
from collections import defaultdict

import httpx
import openpyxl
import psycopg2
from psycopg2.extras import execute_values

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger("acordofes_ingest")

PAGINA = "https://www.saude.mg.gov.br/acordofes"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0) Chrome/131"}

# indices das colunas no Excel (aba "VALOR RESIDUAL")
C_DIV_INI, C_PAGO, C_DIV_ATUAL, C_RETIRADO, C_PAGO_FORA, C_CNPJ, C_RAZAO = 12, 19, 20, 21, 22, 24, 25
_PREFIXO_FMS = "FUNDO MUNICIPAL DE SAUDE DE "


def _sync_url() -> str:
    return (os.getenv("DATABASE_URL_SYNC", "")
            .replace("&channel_binding=require", "").replace("?channel_binding=require", ""))


def _norm(s) -> str:
    s = "".join(c for c in unicodedata.normalize("NFKD", str(s or "")) if not unicodedata.combining(c))
    return " ".join(s.upper().split())


def _f(x) -> float:
    try:
        return float(x or 0)
    except (TypeError, ValueError):
        return 0.0


def _xlsx_url() -> str:
    with httpx.Client(timeout=60, verify=False, follow_redirects=True, headers=UA) as c:
        html = c.get(PAGINA).text
    # link do espelho de valores (robusto a mudanca da data no caminho)
    for m in re.finditer(r'href="([^"]+\.xlsx)"', html, re.I):
        if "valores-acordo" in m.group(1).lower():
            return m.group(1)
    raise RuntimeError("link do Valores-acordo.xlsx nao encontrado na pagina do Acordo FES")


C_ANO_EMP, C_EMPENHO, C_RESOLUCAO = 1, 2, 9


def _empenho(r) -> tuple | None:
    """(ano, nº do empenho, resolução, dívida inicial, pago, dívida atual, retirado,
    pago fora) de uma linha da planilha — a chave que o `ses_mg_resolucoes` usa para
    dizer qual resto a pagar pago pela SES é de empenho que está no Acordo."""
    try:
        ano = int(str(r[C_ANO_EMP]).strip())
        emp = str(int(float(str(r[C_EMPENHO]).strip())))
    except (TypeError, ValueError):
        return None
    res = str(r[C_RESOLUCAO] or "").strip()[:20] or None
    return (ano, emp, res, _f(r[C_DIV_INI]), _f(r[C_PAGO]), _f(r[C_DIV_ATUAL]),
            _f(r[C_RETIRADO]), _f(r[C_PAGO_FORA]))


def _grava_empenhos(cur, casados: list, por_empenho: dict) -> None:
    """`acordofes_empenho`, só dos credores casados a um município do tenant.

    ⚠️ Num SAVEPOINT: a tabela é de `add_ses_mg_resolucoes.sql`, e um banco onde ela
    falhou não pode perder o agregado `acordofes_credor`, que é a tela do Acordo."""
    linhas = [(mid, cnpj[:14], *e) for cnpj, mid in casados for e in por_empenho.get(cnpj[:14], [])]
    cur.execute("SAVEPOINT acordofes_empenho")
    try:
        cur.execute("TRUNCATE acordofes_empenho RESTART IDENTITY")
        if linhas:
            execute_values(cur, """INSERT INTO acordofes_empenho
                (municipio_id, cnpj, ano_empenho, num_empenho, resolucao, divida_inicial,
                 total_pago, divida_atual, valor_retirado, pago_fora) VALUES %s""", linhas)
        cur.execute("RELEASE SAVEPOINT acordofes_empenho")
        log.info(f"Acordo FES: {len(linhas)} empenho(s) dos credores casados.")
    except Exception as e:
        cur.execute("ROLLBACK TO SAVEPOINT acordofes_empenho")
        log.warning(f"acordofes_empenho nao gravado ({type(e).__name__}); o agregado segue")


try:
    from ingestion import status_coleta as _st
except ImportError:
    import status_coleta as _st


def ingest() -> int:
    # ⚠️ SO MUNICIPIO DE MINAS. O Acordo FES e um programa do Estado de MG, e o
    # match e por NOME (razao social do Fundo de Saude): sem o filtro de UF, a
    # divida de uma cidade mineira colava num HOMONIMO de outro estado. O
    # cadastro vem ANTES do download — tenant sem MG nem baixa a planilha.
    conn = psycopg2.connect(_sync_url())
    cur = conn.cursor()
    cur.execute("SELECT id, nome FROM municipios WHERE upper(coalesce(uf, '')) = 'MG'")
    mun_by_nome = {_norm(n): i for i, n in cur.fetchall()}
    if not mun_by_nome:
        log.info("nenhum municipio de MG neste tenant — Acordo FES nao se aplica; nada baixado")
        conn.close()
        return 0

    url = _xlsx_url()
    log.info(f"baixando {url.split('/')[-1]} ...")
    with httpx.Client(timeout=180, verify=False, follow_redirects=True, headers=UA) as c:
        content = c.get(url).content
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    # agrega por CNPJ
    agg = {}  # cnpj -> [div_ini, pago, div_atual, retirado, pago_fora, n, razao]
    por_empenho = defaultdict(list)  # cnpj -> [(ano, empenho, resolucao, valores...)]
    rows = ws.iter_rows(values_only=True)
    next(rows, None)  # cabecalho
    for r in rows:
        if not r or len(r) <= C_RAZAO:
            continue
        cnpj = re.sub(r"\D", "", str(r[C_CNPJ] or ""))
        if not cnpj:
            continue
        emp = _empenho(r)
        if emp:
            por_empenho[cnpj[:14]].append(emp)
        a = agg.get(cnpj)
        if a is None:
            a = agg[cnpj] = [0.0, 0.0, 0.0, 0.0, 0.0, 0, str(r[C_RAZAO] or "").strip()]
        a[0] += _f(r[C_DIV_INI]); a[1] += _f(r[C_PAGO]); a[2] += _f(r[C_DIV_ATUAL])
        a[3] += _f(r[C_RETIRADO]); a[4] += _f(r[C_PAGO_FORA]); a[5] += 1

    batch = []
    matched = 0
    for cnpj, a in agg.items():
        razao = a[6]
        rn = _norm(razao)
        mid = None
        if rn.startswith(_PREFIXO_FMS):
            mid = mun_by_nome.get(rn[len(_PREFIXO_FMS):].strip())
            if mid:
                matched += 1
        batch.append((cnpj[:14], razao, mid, a[0], a[1], a[2], a[3], a[4], a[5]))

    cur.execute("TRUNCATE acordofes_credor RESTART IDENTITY")
    execute_values(cur, """INSERT INTO acordofes_credor
        (cnpj, razao_social, municipio_id, divida_inicial, total_pago,
         divida_atual, valor_retirado, pago_fora, n_empenhos) VALUES %s""", batch)
    _grava_empenhos(cur, [(b[0], b[2]) for b in batch if b[2]], por_empenho)
    conn.commit()
    # ⚠️ A-1 (auditoria 11/09): status HONESTO. `batch`=0 => a planilha veio vazia /
    # o parse (indices de coluna fixos) quebrou => error (implausivel 0 credores).
    # `matched`=0 com credores carregados => partial (o dado nacional entrou, mas
    # nenhum "FUNDO MUNICIPAL DE SAUDE DE ..." casou um municipio — vinculo quebrou).
    status, erro = _st.acordofes(len(batch), matched)
    try:
        cur.execute("INSERT INTO ingestion_log (source, status, records_inserted, error_message, finished_at) "
                    "VALUES ('acordofes',%s,%s,%s,NOW())", (status, len(batch), erro))
        conn.commit()
    except Exception:
        conn.rollback()
    cur.close()
    conn.close()
    log.info(f"Acordo FES: {len(batch)} credores ({matched} casados aos nossos municipios).")
    return len(batch)


if __name__ == "__main__":
    ingest()
